"""Build anonymous official-template competition documents from versioned prose."""
import argparse
import copy
import hashlib
import json
import math
import re
import zipfile
from fractions import Fraction
from pathlib import Path

import pypdfium2 as pdfium
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.graphics import renderPDF
from reportlab.graphics.shapes import Drawing, String
from reportlab.graphics.charts.lineplots import LinePlot
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.charts.legends import Legend
from reportlab.graphics.widgets.markers import makeMarker
from reportlab.lib import colors
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
TASK = ROOT / ".build/submission-20261018"
NUMBER = "CACR2026E7YZBA"
TITLE = "面向新型电网的细粒度密态聚合与可信数据协同平台"
REFERENCE = next((ROOT / "references/competition/official-attachments/expanded").rglob("*模板*.docx"))
APPROVAL = next((ROOT / "references/competition/official-attachments/expanded").rglob("附件三*.docx"))
FONT = "C:/Windows/Fonts/msyh.ttc"
BLUE, TEAL, GRAY = "#153a52", "#157968", "#62798a"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def ratio(item):
    return Fraction(int(item["numerator"]), int(item["denominator"]))


def replace(paragraph, value):
    texts = paragraph._p.xpath(".//w:t")
    if texts:
        texts[0].text = value
        for text in texts[1:]:
            text.text = ""
    else:
        paragraph.add_run(value)


def figures(directory):
    directory.mkdir(parents=True, exist_ok=True)
    font = ImageFont.truetype(FONT, 33)
    small = ImageFont.truetype(FONT, 28)
    title = ImageFont.truetype(FONT, 42)

    def base():
        image = Image.new("RGB", (1800, 920), "white")
        return image, ImageDraw.Draw(image)

    def box(draw, area, heading, lines, fill="#f1f6fa", border="#b6c9d7"):
        draw.rounded_rectangle(area, radius=12, fill=fill, outline=border, width=2)
        x, y, _, _ = area
        draw.text((x + 25, y + 24), heading, font=title, fill=BLUE)
        for i, line in enumerate(lines):
            draw.text((x + 25, y + 95 + i * 49), line, font=font, fill=BLUE)

    def arrow(draw, start, end, label=""):
        draw.line((start, end), fill=TEAL, width=5)
        x, y = end
        dx, dy = x-start[0], y-start[1]
        length = math.hypot(dx, dy)
        ux, uy = dx/length, dy/length
        draw.polygon([(x, y), (x-18*ux-9*uy, y-18*uy+9*ux), (x-18*ux+9*uy, y-18*uy-9*ux)], fill=TEAL)
        if label:
            draw.text(((start[0] + end[0]) / 2, start[1] - 49), label, font=small, fill=GRAY, anchor="mm")

    image, draw = base()
    box(draw, (525, 35, 1250, 195), "TA 参数与身份登记", ["公开参数  /  角色证书  /  可信密钥登记"])
    for i in range(3):
        y = 275 + i * 135
        draw.rounded_rectangle((40, y, 420, y + 108), radius=9, outline="#b6c9d7", fill="#f1f6fa", width=2)
        draw.text((65, y + 17), f"Source {i+1}  读数源侧加密", font=font, fill=BLUE)
        draw.text((65, y + 62), "独立身份、密钥与数据分片", font=small, fill=GRAY)
        arrow(draw, (425, y + 54), (635, 485))
    box(draw, (650, 325, 1120, 650), "GW 聚合网关", ["核对任务授权", "联合检查与密文筛选", "密态聚合并签署回执"])
    arrow(draw, (1125, 485), (1295, 485), "聚合密文")
    box(draw, (1310, 325, 1770, 650), "UC 统计端", ["核验回执与恢复统计", "核对分组、完整性与历史", "签署结果与发布决定"], "#edf6f2")
    draw.line((885, 200, 885, 310), fill=GRAY, width=2)
    draw.line((1245, 115, 1540, 115, 1540, 310), fill=GRAY, width=2)
    draw.text((70, 740), "Source → GW 与 GW → UC：实际混合TLS、双向身份核对、业务数据保护", font=font, fill=BLUE)
    draw.text((70, 800), "公开出口：按发布决定提供群体结果、签名声明与关联证据", font=font, fill=TEAL)
    draw.text((70, 861), "架构示意；本机演示采用1—3个同机独立Source进程", font=small, fill=GRAY)
    image.save(directory / "architecture.png")

    image, draw = base()
    box(draw, (40, 70, 650, 280), "登记完整固定分组", ["至少10个独立计量点", "按数据族沿革确定发布域"])
    box(draw, (1070, 70, 1760, 280), "跨任务共享发布历史", ["关联账户、函数与读数版本", "受保护提交状态与恢复核对"])
    arrow(draw, (655, 175), (875, 380))
    draw.line((1400, 285, 1400, 340, 1050, 340), fill=TEAL, width=5)
    box(draw, (520, 390, 1280, 650), "UC 发布检查", ["完整分组  /  登记规模  /  完整上报", "参与归属、重复点与共享历史核对", "四日期评估整体提交"])
    arrow(draw, (1285, 520), (1460, 520))
    box(draw, (1470, 410, 1770, 650), "签名决定", ["发布结果", "或保留决定"], "#edf6f2")
    draw.text((80, 745), "统一公开投影：页面、接口、事件、自检、评估与证据导出", font=title, fill=BLUE)
    draw.text((80, 820), "规则示意；每次公开提供均关联相应的任务、统计声明与发布状态", font=font, fill=GRAY)
    image.save(directory / "publication.png")

    image, draw = base()
    items = [(40, "任务授权", ["范围、数据、时段", "身份、算法、密钥版本"]),
             (630, "逐源密文快照", ["SM3 Merkle根", "批次、报告与包含关系"]),
             (1220, "聚合回执", ["授权与快照关联", "聚合密文、计数、审计"])]
    for x, heading, lines in items:
        box(draw, (x, 70, x + 530, 325), heading, lines)
    arrow(draw, (575, 200), (620, 200)); arrow(draw, (1165, 200), (1210, 200))
    box(draw, (40, 465, 570, 720), "业务指标", ["16份签名输入精确复算", "四日期关系与四项指标"], "#edf6f2")
    box(draw, (630, 465, 1160, 720), "发布决定", ["规则、任务与发布历史", "统一对外结果投影"], "#edf6f2")
    box(draw, (1220, 465, 1750, 720), "统计端结果", ["统计值与回执摘要", "来源、内容与指定任务"])
    draw.line((1485, 330, 1485, 445), fill=TEAL, width=5)
    arrow(draw, (1210, 590), (1175, 590)); arrow(draw, (620, 590), (585, 590))
    draw.text((65, 805), "签名受理规则：ECDSA与ML-DSA-65同时有效；国密增强另要求SM2", font=font, fill=BLUE)
    draw.text((65, 862), "证据关联示意；签名、可信登记和业务上下文分别核验", font=small, fill=GRAY)
    image.save(directory / "evidence.png")

    pdfmetrics.registerFont(TTFont("GridYahei", FONT, subfontIndex=0))
    def chart_png(drawing, name):
        pdf = directory / (name+".pdf")
        renderPDF.drawToFile(drawing, str(pdf))
        document = pdfium.PdfDocument(str(pdf))
        document[0].render(scale=220/72).to_pil().save(directory / (name+".png"))
        document.close()

    def legend(drawing, pairs):
        item = Legend(); item.x = 82; item.y = 325; item.fontName = "GridYahei"; item.fontSize = 12
        item.columnMaximum = 1; item.alignment = "right"; item.dxTextSpace = 10; item.deltax = 170
        item.colorNamePairs = [(colors.HexColor(color), label) for color, label in pairs]
        drawing.add(item)
    evaluation_path = ROOT / "artifacts/runs/67027b12c5884f9f8f3ce99415bc58c3/evaluation.json"
    evaluation = read(evaluation_path)
    curve = evaluation["metrics"]["curve"]
    baseline = [float(ratio(x["baseline_W"])) / 1000 for x in curve]
    event = [x["event_W"] / 1000 for x in curve]
    drawing = Drawing(760, 380)
    plot = LinePlot(); plot.x = 80; plot.y = 85; plot.width = 620; plot.height = 205
    plot.data = [list(enumerate(baseline)), list(enumerate(event))]
    plot.xValueAxis.valueMin = 0; plot.xValueAxis.valueMax = 3; plot.xValueAxis.valueSteps = list(range(4))
    labels = [x["slot"]["slot_start"][11:16] for x in curve]
    plot.xValueAxis.labelTextFormat = lambda value: labels[int(value)]
    plot.xValueAxis.labels.fontName = "GridYahei"; plot.xValueAxis.labels.fontSize = 12
    plot.yValueAxis.labels.fontName = "GridYahei"; plot.yValueAxis.labels.fontSize = 12
    plot.yValueAxis.valueMin = 1350; plot.yValueAxis.valueMax = 1650; plot.yValueAxis.valueStep = 75
    plot.lines[0].strokeColor = colors.HexColor("#72879a"); plot.lines[0].strokeDashArray = [5, 4]
    plot.lines[1].strokeColor = colors.HexColor(TEAL)
    for index in range(2):
        line = plot.lines[index]
        line.strokeWidth = 2; line.symbol = makeMarker("FilledCircle"); line.symbol.size = 6
    drawing.add(plot)
    drawing.add(String(80, 355, "300输入 / 150授权参与点 / 合成场景预设10%响应", fontName="GridYahei", fontSize=14, fillColor=colors.HexColor(BLUE)))
    drawing.add(String(80, 303, "群体功率（kW）", fontName="GridYahei", fontSize=12))
    drawing.add(String(280, 34, "事件窗口起始时刻（UTC）", fontName="GridYahei", fontSize=12))
    legend(drawing, [("#72879a", "三日参考负荷"), (TEAL, "事件日负荷")])
    chart_png(drawing, "response")

    summary_path = ROOT / "artifacts/iterations/privacy/benchmarks/current.json"
    summary = read(summary_path)
    rows = summary["configurations"]
    drawing = Drawing(760, 390)
    plot = VerticalBarChart(); plot.x = 80; plot.y = 95; plot.width = 620; plot.height = 215
    plot.data = [[float(x["arms"][arm]["median_elapsed_ns"]) / 1e9 for x in rows] for arm in ["C0", "C1"]]
    plot.categoryAxis.categoryNames = [f'{x["nodes"]}输入 / {x["registered_metering_points"]}授权\n{x["formal_pairs"]}对 / {x["core_source_digest"][:8]}' for x in rows]
    plot.categoryAxis.labels.fontName = "GridYahei"; plot.categoryAxis.labels.fontSize = 11
    plot.categoryAxis.labels.boxAnchor = "n"; plot.categoryAxis.labels.dy = -8
    plot.valueAxis.valueMin = 0; plot.valueAxis.valueMax = 270; plot.valueAxis.valueStep = 50
    plot.valueAxis.labels.fontName = "GridYahei"; plot.valueAxis.labels.fontSize = 12
    plot.bars[0].fillColor = colors.HexColor("#a5b7c5"); plot.bars[1].fillColor = colors.HexColor(TEAL)
    plot.bars.strokeColor = None; plot.groupSpacing = 28; plot.barSpacing = 3
    plot.barLabelFormat = "%.2f"; plot.barLabels.fontName = "GridYahei"; plot.barLabels.fontSize = 10; plot.barLabels.nudge = 6
    drawing.add(plot)
    drawing.add(String(80, 365, "四日期评估耗时中位数（秒）", fontName="GridYahei", fontSize=14, fillColor=colors.HexColor(BLUE)))
    legend(drawing, [("#a5b7c5", "C0 密态实验对照"), (TEAL, "C1 启用发布保护")])
    chart_png(drawing, "performance")
    write_json(directory / "figure-inputs.json", {"evaluation": {"path": str(evaluation_path.relative_to(ROOT)), "sha256": digest(evaluation_path)},
                                                "benchmarks": {"path": str(summary_path.relative_to(ROOT)), "sha256": digest(summary_path)},
                                                "chart_scope": "Real stored observations; architecture/publication/evidence images are labelled schematics"})


def make_math_text(text):
    run = OxmlElement("m:r")
    value = OxmlElement("m:t"); value.text = text; run.append(value)
    return run


def subscript(base, sub):
    node = OxmlElement("m:sSub")
    base_node = OxmlElement("m:e"); base_node.append(make_math_text(base)); node.append(base_node)
    sub_node = OxmlElement("m:sub"); sub_node.append(make_math_text(sub)); node.append(sub_node)
    return node


def fraction(top, bottom):
    node = OxmlElement("m:f")
    for key, items in [("num", top), ("den", bottom)]:
        container = OxmlElement("m:" + key)
        for item in items:
            container.append(make_math_text(item) if isinstance(item, str) else item)
        node.append(container)
    return node


def sum_term(lower, upper, expression):
    node = OxmlElement("m:nary"); props = OxmlElement("m:naryPr")
    char = OxmlElement("m:chr"); char.set(qn("m:val"), "∑"); props.append(char)
    loc = OxmlElement("m:limLoc"); loc.set(qn("m:val"), "subSup"); props.append(loc); node.append(props)
    for key, value in [("sub", [lower]), ("sup", [upper]), ("e", expression)]:
        part = OxmlElement("m:" + key)
        for item in value:
            part.append(make_math_text(item) if isinstance(item, str) else item)
        node.append(part)
    return node


def package_with_preserved_parts(reference, generated, target, whitelist):
    with zipfile.ZipFile(reference) as source, zipfile.ZipFile(generated) as draft, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as output:
        for item in draft.infolist():
            data = source.read(item.filename) if item.filename in source.namelist() and item.filename not in whitelist else draft.read(item.filename)
            if item.filename == "docProps/app.xml":
                tree = etree.fromstring(data)
                for element in tree:
                    if etree.QName(element).localname in ("Company", "Manager"):
                        element.text = None
                data = etree.tostring(tree, encoding="UTF-8", xml_declaration=True, standalone=True)
            output.writestr(item.filename, data)
        for item in source.infolist():
            if item.filename not in draft.namelist():
                output.writestr(item, source.read(item))
    with zipfile.ZipFile(reference) as source, zipfile.ZipFile(target) as result:
        preserved = [name for name in source.namelist() if name not in whitelist]
        assert all(source.read(name) == result.read(name) for name in preserved)
    return preserved


def build(source_path, destination, abstract):
    source_path = source_path.resolve()
    destination = destination.resolve()
    assert digest(REFERENCE) == "4f81348bf085527f7ebe5de2787fd91ba05e5b5531754916badb005fcfb512c2"
    destination.mkdir(parents=True, exist_ok=True)
    assets = TASK / "figures"; figures(assets)
    doc = Document(REFERENCE); original = list(doc.paragraphs)
    components = {1: copy.deepcopy(original[27]._p), 2: copy.deepcopy(original[31]._p), 0: copy.deepcopy(original[32]._p)}
    replace(original[0], "作品编号：" + NUMBER)
    for run in original[0].runs:
        if run.text:
            from docx.shared import RGBColor
            run.font.color.rgb = RGBColor(0, 0, 0)
    replace(original[3], "☑密码应用技术  □其它")
    replace(original[14], "作品题目：" + TITLE)
    replace(original[17], "B类自拟选题")
    replace(original[22], "2026年10月18日")
    basic = [row.cells[0] for row in doc.tables[0].rows]
    info = ["编号："+NUMBER, "作品题目："+TITLE,
            "作品类别：□软件设计  □硬件制作  □工程实践  ☑密码技术应用  □其它",
            "作品内容摘要：\n"+abstract, "关键词（五个）：细粒度密态聚合  PFDAM  受控发布  联合签名  能源数据协同"]
    for cell, value in zip(basic[1:], info):
        replace(cell.paragraphs[0], value)
        for paragraph in list(cell.paragraphs[1:]):
            paragraph._p.getparent().remove(paragraph._p)
    replace(basic[3].paragraphs[0], "作品类别：")
    category = basic[3].paragraphs[0].add_run("□软件设计  □硬件制作  □工程实践  ☑密码技术应用  □其它")
    category.font.name = "楷体_GB2312"; category.font.size = Pt(12)
    category._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "楷体_GB2312")
    for paragraph in original[27:]:
        if not paragraph._p.xpath(".//w:sectPr"):
            paragraph._p.getparent().remove(paragraph._p)

    def paragraph(value="", level=0, secondary=False):
        node = copy.deepcopy(components[level])
        formatting = node.xpath("./w:r/w:rPr")
        formatting = copy.deepcopy(formatting[0]) if formatting else None
        for child in list(node):
            if child.tag != qn("w:pPr"):
                node.remove(child)
        run = OxmlElement("w:r")
        if formatting is not None:
            run.append(formatting)
        text = OxmlElement("w:t"); text.text = value; run.append(text); node.append(run)
        doc._body._body.insert(len(doc._body._body)-1, node)
        p = Paragraph(node, doc._body)
        if level:
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
            p.paragraph_format.space_after = Pt(6)
        else:
            p.paragraph_format.first_line_indent = Pt(28)
            p.paragraph_format.space_after = Pt(4)
        if secondary:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.keep_together = True
            for run in p.runs:
                run.font.size = Pt(11)
        return p

    def figure(name, caption):
        p = paragraph(); p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(str(assets / (name+".png")), width=Cm(16.5))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = paragraph(caption); p.paragraph_format.first_line_indent = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(12)

    def table(rows):
        count = len(rows[0]); table = doc.add_table(rows=0, cols=count)
        table.autofit = False
        widths = {3: [3.6, 7.3, 5.5], 4: [3.7, 6.0, 4.0, 2.7], 6: [2.4, 1.6, 2.4, 2.4, 4.2, 3.4]}.get(count, [16.4/count]*count)
        total = sum(widths); widths = [x*16.4/total for x in widths]
        for col, width in zip(table.columns, widths): col.width = Cm(width)
        props = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for name in ("top", "bottom", "left", "right", "insideH", "insideV"):
            edge = OxmlElement("w:"+name); edge.set(qn("w:val"), "single"); edge.set(qn("w:sz"), "4"); edge.set(qn("w:color"), "BFC9D0"); borders.append(edge)
        props.append(borders)
        for index, values in enumerate(rows):
            row = table.add_row(); row_properties = row._tr.get_or_add_trPr()
            row_properties.append(OxmlElement("w:cantSplit"))
            if index == 0: row_properties.append(OxmlElement("w:tblHeader"))
            for cell, value, width in zip(row.cells, values, widths):
                cell.width = Cm(width)
                p = cell.paragraphs[0]; p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.15
                if len(rows) <= 5 and index < len(rows)-1:
                    p.paragraph_format.keep_with_next = True
                run = p.add_run(value); run.font.name = "楷体_GB2312"; run.font.size = Pt(12); run.bold = index == 0
                run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "楷体_GB2312")
                margins = OxmlElement("w:tcMar")
                for side in ("top", "bottom", "left", "right"):
                    margin = OxmlElement("w:"+side); margin.set(qn("w:w"), "80"); margin.set(qn("w:type"), "dxa"); margins.append(margin)
                cell._tc.get_or_add_tcPr().append(margins)
        paragraph()

    def equation(kind):
        if kind == "aggregate":
            sequences = [["M = ", sum_term("i∈A", "", [subscript("m", "i")]), "， AVG(A) = ", fraction(["M"], ["|A|"])]]
        else:
            delta = ["(", subscript("B", "t"), " − ", subscript("P", "e,t"), ")"]
            sequences = [[subscript("B", "t"), " = ", fraction([subscript("P", "d1,t"), "+", subscript("P", "d2,t"), "+", subscript("P", "d3,t")], ["3"])],
                         [subscript("ΔP", "mean"), " = ", fraction(["1"], ["4"]), sum_term("t=1", "4", copy.deepcopy(delta))],
                         ["ΔE = ", sum_term("t=1", "4", copy.deepcopy(delta)), " × 0.25 h"]]
        for sequence in sequences:
            p = paragraph(); p.paragraph_format.first_line_indent = Pt(0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            mathpara = OxmlElement("m:oMathPara"); mathnode = OxmlElement("m:oMath")
            for item in sequence: mathnode.append(make_math_text(item) if isinstance(item, str) else item)
            mathpara.append(mathnode); p._p.append(mathpara)

    chunks = source_path.read_text(encoding="utf-8").strip().split("\n\n")
    for chunk in chunks:
        if chunk.startswith("# "): continue
        if chunk.startswith("## "): paragraph(chunk[3:], 1)
        elif chunk.startswith("### "): paragraph(chunk[4:], 2)
        elif chunk.startswith("{{figure:"):
            name, caption = chunk[len("{{figure:"):-2].split("|", 1); figure(name, caption)
        elif chunk.startswith("{{equation:"): equation(chunk[len("{{equation:"):-2])
        elif chunk.startswith("|"):
            rows = [[x.strip() for x in line.strip().strip("|").split("|")] for line in chunk.splitlines()]
            table([row for row in rows if not all(re.fullmatch(r"[:\- ]+", cell) for cell in row)])
        else: paragraph(chunk.replace("\n", " "), secondary=bool(re.match(r"\[\d+\]", chunk)))

    doc.core_properties.author = ""; doc.core_properties.last_modified_by = ""; doc.core_properties.title = TITLE
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None: update = OxmlElement("w:updateFields"); settings.append(update)
    update.set(qn("w:val"), "true")
    working = destination / "report-working.docx"; doc.save(working)
    target = destination / f"{NUMBER}-2026密码竞赛复赛作品设计报告.docx"
    whitelist = {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml", "docProps/core.xml", "docProps/app.xml", "word/settings.xml"}
    preserved = package_with_preserved_parts(REFERENCE, working, target, whitelist)
    assert len(Document(target).sections) == 3
    output_approval = destination / f"{NUMBER}-2026密码竞赛审批及签字页_待签署.docx"
    approval = Document(APPROVAL)
    seen = set()
    for tbl in approval.tables:
        for row in tbl.rows:
            for cell in row.cells:
                if cell._tc in seen: continue
                seen.add(cell._tc)
                for p in cell.paragraphs:
                    if p.text.startswith("作品编号："): replace(p, "作品编号："+NUMBER)
                    elif p.text.startswith("作品题目："): replace(p, "作品题目："+TITLE)
    approval.core_properties.author = ""; approval.core_properties.last_modified_by = ""
    # Word requires a paragraph after the full-page form. Keep that empty
    # paragraph on the form page without changing any field or table geometry.
    for p in approval.paragraphs:
        if not p.text and not p._p.xpath('.//w:sectPr'):
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(1)
            props = p._p.get_or_add_pPr()
            rprops = props.find(qn('w:rPr'))
            if rprops is None:
                rprops = OxmlElement('w:rPr'); props.append(rprops)
            size = rprops.find(qn('w:sz'))
            if size is None:
                size = OxmlElement('w:sz'); rprops.append(size)
            size.set(qn('w:val'), '2')
    approval_working = destination / "approval-working.docx"; approval.save(approval_working)
    approval_preserved = package_with_preserved_parts(APPROVAL, approval_working, output_approval,
                                                    {"word/document.xml", "docProps/core.xml", "docProps/app.xml"})
    write_json(destination / "fidelity.json", {"template_sha256": digest(REFERENCE), "report_sha256": digest(target),
               "source_sha256": digest(source_path), "preserved_parts": preserved, "sections": 3,
               "approval_source_sha256": digest(APPROVAL), "approval_sha256": digest(output_approval),
               "approval_preserved_parts": approval_preserved, "approval_signatures_added": False,
               "source": str(source_path.relative_to(ROOT)), "outputs": [target.name, output_approval.name]})
    print(json.dumps({"report": str(target), "approval": str(output_approval), "preserved_parts": len(preserved)}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(); parser.add_argument("--source", required=True); parser.add_argument("--out", required=True); parser.add_argument("--abstract", required=True)
    args = parser.parse_args()
    build(Path(args.source), Path(args.out), Path(args.abstract).read_text(encoding="utf-8").strip())
